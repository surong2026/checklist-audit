from docx import Document


def parse_docx(file_path: str) -> dict:
    """Extract all tables from a .docx file. Returns the largest table as primary."""
    doc = Document(file_path)
    tables_data = []

    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
        if rows:
            tables_data.append({"index": ti, "rows": rows, "header": rows[0] if rows else []})

    if not tables_data:
        # Try extracting from paragraphs as fallback
        all_text = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                all_text.append(t)
        return {"tables": [], "paragraphs": all_text, "primary_table": None}

    # Pick the table with the most rows as primary data table
    primary = max(tables_data, key=lambda t: len(t["rows"]))

    return {
        "tables": tables_data,
        "paragraphs": [],
        "primary_table": primary["rows"],
        "primary_index": primary["index"],
    }
