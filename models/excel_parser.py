import openpyxl
import xlrd


def parse_xlsx(file_path: str) -> dict:
    """Parse .xlsx file with openpyxl. Returns first sheet as primary data."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(min_row=1, values_only=True):
            rows.append([str(c) if c is not None else "" for c in row])
        sheets[name] = rows

    primary = sheets[wb.sheetnames[0]] if sheets else []
    return {"sheets": sheets, "primary_table": primary, "primary_sheet": wb.sheetnames[0] if sheets else None}


def parse_xls(file_path: str) -> dict:
    """Parse .xls file with xlrd."""
    wb = xlrd.open_workbook(file_path)
    sheets = {}
    for name in wb.sheet_names():
        ws = wb.sheet_by_name(name)
        rows = []
        for r in range(ws.nrows):
            rows.append([str(ws.cell_value(r, c)) if ws.cell_value(r, c) != "" else "" for c in range(ws.ncols)])
        sheets[name] = rows

    primary = sheets[wb.sheet_names()[0]] if sheets else []
    return {"sheets": sheets, "primary_table": primary, "primary_sheet": wb.sheet_names()[0] if sheets else None}
